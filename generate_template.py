#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 7班 班级课表 Excel 模板（供用户填写后导入应用）"""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

OUT = "/Users/a123/WorkBuddy/2026-09-07-11-19-53/7班课表模板.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "7班课表"

# 样式
header_font = Font(bold=True, size=12)
group_font = Font(bold=True, size=11)
center = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = PatternFill("solid", fgColor="F2F6F5")

headers = ["节次", "星期1", "星期2", "星期3", "星期4", "星期5"]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = header_font
    cell.alignment = center
    cell.border = border

# 每组时段（顺序：上午→下午→晚自习，无大课间操/中午休息/眼保健健康等分隔行）
groups = [
    ("上午", ["一", "二", "三", "四", "五"]),
    ("下午", ["六", "七", "八", "九"]),
    ("晚自习", ["1", "2", "3", "4"]),
]

r = 2
for gi, (gname, periods) in enumerate(groups):
    # 分组标题行
    gcell = ws.cell(row=r, column=1, value=gname)
    gcell.font = group_font
    gcell.alignment = center
    gcell.border = border
    r += 1
    for p in periods:
        pcell = ws.cell(row=r, column=1, value=p)
        pcell.alignment = center
        pcell.border = border
        pcell.fill = fill
        for c in range(2, 7):
            dcell = ws.cell(row=r, column=c, value="")
            dcell.alignment = center
            dcell.border = border
        r += 1

# 列宽
widths = [10, 14, 14, 14, 14, 14]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

wb.save(OUT)
print("已生成模板:", OUT)
print("行数:", r - 1)
